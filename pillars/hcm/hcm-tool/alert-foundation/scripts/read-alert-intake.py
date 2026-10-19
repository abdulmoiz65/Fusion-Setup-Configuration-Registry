#!/usr/bin/env python3
"""Read an Alert Foundation XLSX intake and produce a REST payload plan.

The script is intentionally local and read-only. It never calls Fusion and
never treats workbook status or approval cells as execution authorization.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any
from urllib.parse import quote, urlsplit

sys.path.insert(0, str(Path(__file__).resolve().parent))
import alert_validation


RESOURCE = "/hcmRestApi/resources/11.13.18.05/eventAlerts"
CONTRACT_VERSION = alert_validation.VERSION
HEADER_ROW = 4
DATA_START_ROW = 5
ALERT_FIELDS = (
    "AlertCode",
    "Name",
    "Description",
    "AlertType",
    "AlertAccessLevel",
    "EnabledFlag",
    "TestModeFlag",
)
TEMPLATE_FIELDS = (
    "TemplateCode",
    "Name",
    "MimeType",
    "Subject",
    "DefaultLanguage",
    "EnabledFlag",
)
RECIPIENT_FIELDS = (
    "RecipientCode",
    "ChannelType",
    "ValueExpression",
)
BOOLEAN_FIELDS = {"EnabledFlag", "TestModeFlag"}


class IntakeError(ValueError):
    """A caller-correctable workbook or argument error."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def canonical_sha256(value: Any) -> str:
    rendered = json.dumps(
        value, sort_keys=True, separators=(",", ":"), ensure_ascii=False
    ).encode("utf-8")
    return hashlib.sha256(rendered).hexdigest().upper()


def normalize_target_environment(value: str | None) -> str | None:
    if value is None or not value.strip():
        return None
    parsed = urlsplit(value.strip())
    if (
        parsed.scheme.lower() != "https"
        or not parsed.netloc
        or parsed.query
        or parsed.fragment
        or parsed.path.rstrip("/")
    ):
        raise IntakeError(
            "target environment must be an HTTPS Fusion origin without a path, "
            "query, or fragment"
        )
    return f"https://{parsed.netloc}"


def normalize_header(value: Any) -> str:
    text = str(value or "").strip()
    return text[:-1] if text.endswith("*") else text


def normalize_bool(value: Any, *, location: str) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip().upper() in {"TRUE", "FALSE"}:
        return value.strip().upper() == "TRUE"
    raise IntakeError(f"{location}: expected TRUE or FALSE")


def required(row: dict[str, Any], field: str, *, location: str) -> Any:
    value = row.get(field)
    if value is None or (isinstance(value, str) and not value.strip()):
        raise IntakeError(f"{location}: missing required field {field}")
    if isinstance(value, str) and value.strip().upper().startswith("REPLACE_WITH_"):
        raise IntakeError(f"{location}: unresolved placeholder in {field}")
    return value


def parse_mapping(values: list[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for item in values:
        if "=" not in item:
            raise IntakeError(f"Expected source=rest-code mapping, got {item!r}")
        source, target = (part.strip() for part in item.split("=", 1))
        if not source or not target:
            raise IntakeError(f"Expected source=rest-code mapping, got {item!r}")
        result[source.lower()] = target
    return result


def open_workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError(
            "openpyxl is required for the basic XLSX intake backend"
        ) from exc

    if path.suffix.lower() != ".xlsx":
        raise IntakeError("The basic backend accepts .xlsx files only")
    if not path.is_file():
        raise IntakeError(f"Workbook does not exist: {path}")
    try:
        values = load_workbook(
            path, read_only=True, data_only=True, keep_links=False
        )
        formulas = load_workbook(
            path, read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise IntakeError(f"Unable to read XLSX workbook: {exc}") from exc
    return values, formulas


def read_rows(values_wb, formulas_wb, sheet_name: str) -> list[dict[str, Any]]:
    from openpyxl.utils.cell import get_column_letter

    if sheet_name not in values_wb.sheetnames:
        raise IntakeError(f"Missing required sheet {sheet_name!r}")
    values_ws = values_wb[sheet_name]
    formulas_ws = formulas_wb[sheet_name]
    headers = [normalize_header(cell.value) for cell in values_ws[HEADER_ROW]]
    if not headers or not headers[0]:
        raise IntakeError(f"{sheet_name}!{HEADER_ROW}: missing header row")
    nonblank = [header for header in headers if header]
    if len(nonblank) != len(set(nonblank)):
        raise IntakeError(f"{sheet_name}!{HEADER_ROW}: duplicate headers")

    rows: list[dict[str, Any]] = []
    for row_number in range(DATA_START_ROW, values_ws.max_row + 1):
        value_cells = values_ws[row_number]
        formula_cells = formulas_ws[row_number]
        row = {
            header: value_cells[index].value
            for index, header in enumerate(headers)
            if header
        }
        include = str(row.get("IncludeFlag") or "").strip().upper()
        if include not in {"YES", "NO"}:
            continue
        if include == "NO":
            continue

        for index, header in enumerate(headers):
            if not header or header in {"ValidationStatus", "ValidationMessage"}:
                continue
            raw_formula = formula_cells[index].value
            if isinstance(raw_formula, str) and raw_formula.startswith("="):
                coordinate = formula_cells[index].coordinate
                raise IntakeError(
                    f"{sheet_name}!{coordinate}: formulas are not allowed in intake fields"
                )
        row["_source"] = {
            "sheet": sheet_name,
            "row": row_number,
            "range": f"A{row_number}:{get_column_letter(len(headers))}{row_number}",
        }
        rows.append(row)
    return rows


def load_current_alerts(path: Path | None) -> list[dict[str, Any]]:
    if path is None:
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise IntakeError(f"Unable to read current-state JSON: {exc}") from exc
    if isinstance(data, dict) and isinstance(data.get("items"), list):
        return [item for item in data["items"] if isinstance(item, dict)]
    if isinstance(data, dict) and "AlertCode" in data:
        return [data]
    if isinstance(data, list):
        return [item for item in data if isinstance(item, dict)]
    raise IntakeError("Current-state JSON must be an alert object, items response, or list")


def find_current_alert(
    row: dict[str, Any], current_alerts: list[dict[str, Any]], *, location: str
) -> dict[str, Any]:
    alert_id = str(row.get("AlertId") or "").strip()
    alert_code = str(row.get("AlertCode") or "").strip()
    matches = [
        alert
        for alert in current_alerts
        if (alert_id and str(alert.get("AlertId") or "") == alert_id)
        or (not alert_id and alert_code and alert.get("AlertCode") == alert_code)
    ]
    if len(matches) != 1:
        raise IntakeError(
            f"{location}: expected exactly one current-state match, found {len(matches)}"
        )
    return matches[0]


def alert_payload(row: dict[str, Any], *, location: str) -> dict[str, Any]:
    for field in (
        "ClientAlertKey",
        "Operation",
        "AlertCode",
        "Name",
        "AlertType",
        "AlertAccessLevel",
        "EnabledFlag",
        "TestModeFlag",
        "BusinessOwner",
        "ConsumerPillar",
    ):
        required(row, field, location=location)
    if len(str(row["Name"])) > 80:
        raise IntakeError(f"{location}: Name exceeds 80 characters")
    if row.get("Description") is not None and len(str(row["Description"])) > 2000:
        raise IntakeError(f"{location}: Description exceeds 2000 characters")
    payload = {field: row.get(field) for field in ALERT_FIELDS}
    for field in BOOLEAN_FIELDS:
        payload[field] = normalize_bool(payload[field], location=f"{location}.{field}")
    return payload


def build_alert_operation(
    row: dict[str, Any], current_alerts: list[dict[str, Any]]
) -> dict[str, Any]:
    source = row["_source"]
    location = f"{source['sheet']}!{source['row']}"
    operation = str(required(row, "Operation", location=location)).strip().upper()
    payload = alert_payload(row, location=location)
    client_key = str(row["ClientAlertKey"])

    if operation == "CREATE":
        if row.get("AlertId") not in {None, ""}:
            raise IntakeError(f"{location}: AlertId must be blank for CREATE")
        return {
            "entity": "alert",
            "client_alert_key": client_key,
            "operation": operation,
            "method": "POST",
            "relative_uri": RESOURCE,
            "payload": payload,
            "source": source,
            "preflight": [
                "compatible target describe",
                "zero exact AlertCode matches",
                "live lookup validation",
                "immutable preview and immediate approval",
            ],
        }

    if operation not in {"UPDATE", "DISABLE"}:
        raise IntakeError(
            f"{location}: basic reader supports CREATE, UPDATE, and DISABLE only"
        )
    current = find_current_alert(row, current_alerts, location=location)
    canonical_key = current.get("@context", {}).get("key") or current.get("AlertCode")
    if not canonical_key:
        raise IntakeError(f"{location}: current state has no canonical key")
    if operation == "DISABLE":
        patch = {"EnabledFlag": False}
    else:
        patch = {
            field: value
            for field, value in payload.items()
            if current.get(field) != value
        }
    if not patch:
        raise IntakeError(f"{location}: UPDATE produces no changed fields")
    return {
        "entity": "alert",
        "client_alert_key": client_key,
        "operation": operation,
        "method": "PATCH",
        "relative_uri": f"{RESOURCE}/{quote(str(canonical_key), safe='')}",
        "payload": patch,
        "matched_identity": {
            "AlertId": current.get("AlertId"),
            "AlertCode": current.get("AlertCode"),
            "canonical_key": canonical_key,
        },
        "source": source,
        "preflight": [
            "fresh exact current state",
            "compatible target describe",
            "live lookup validation for changed coded fields",
            "immutable preview and immediate approval",
        ],
    }


def build_template_operation(
    row: dict[str, Any], parent_operations: dict[str, dict[str, Any]], mime_map: dict[str, str],
    validation_mode: str = "deep",
) -> dict[str, Any]:
    source = row["_source"]
    location = f"{source['sheet']}!{source['row']}"
    for field in (
        "ClientAlertKey",
        "TemplateKey",
        "Operation",
        "TemplateCode",
        "Name",
        "MimeType",
        "DefaultLanguage",
        "EnabledFlag",
        "Subject",
    ):
        required(row, field, location=location)
    client_key = str(row["ClientAlertKey"])
    if client_key not in parent_operations:
        raise IntakeError(f"{location}: no included parent for ClientAlertKey {client_key}")
    operation = str(row["Operation"]).strip().upper()
    if operation not in {"CREATE", "UPDATE", "DISABLE"}:
        raise IntakeError(
            f"{location}: basic reader supports CREATE, UPDATE, and DISABLE only"
        )
    mime_source = str(row["MimeType"]).strip()
    mime_type = mime_source if validation_mode == "simple" or mime_source.startswith("ORA_") else mime_map.get(mime_source.lower())
    if not mime_type:
        raise IntakeError(
            f"{location}: MimeType {mime_source!r} requires a live-validated --mime-type-map"
        )
    payload = {
        field: row.get(field)
        for field in TEMPLATE_FIELDS
    }
    payload["MimeType"] = mime_type
    payload["EnabledFlag"] = normalize_bool(
        payload["EnabledFlag"], location=f"{location}.EnabledFlag"
    )
    if len(str(payload["Name"])) > 80:
        raise IntakeError(f"{location}: Name exceeds 80 characters")
    if len(str(payload["Subject"])) > 2000:
        raise IntakeError(f"{location}: Subject exceeds 2000 characters")

    inline_body = row.get("Body")
    has_inline_body = inline_body is not None and inline_body != ""
    if has_inline_body and not isinstance(inline_body, str):
        raise IntakeError(f"{location}: Body must be plain text")

    body_plan = None
    body_path_value = row.get("BodyFilePath")
    body_digest_value = row.get("BodyFileSHA256")
    if has_inline_body and (body_path_value or body_digest_value):
        raise IntakeError(
            f"{location}: use either Body or BodyFilePath/BodyFileSHA256, not both"
        )
    if body_path_value or body_digest_value:
        body_path_text = str(required(row, "BodyFilePath", location=location)).strip()
        expected_digest = str(required(row, "BodyFileSHA256", location=location)).strip().upper()
        body_relative = Path(body_path_text)
        if body_relative.is_absolute() or ".." in body_relative.parts:
            raise IntakeError(f"{location}: BodyFilePath must be safe and relative")
        workbook_dir = Path(row["_workbook_dir"])
        body_path = (workbook_dir / body_relative).resolve()
        if workbook_dir.resolve() not in body_path.parents:
            raise IntakeError(f"{location}: BodyFilePath escapes the workbook directory")
        if not body_path.is_file():
            raise IntakeError(f"{location}: body file does not exist: {body_path_text}")
        actual_digest = sha256_file(body_path)
        if actual_digest != expected_digest:
            raise IntakeError(f"{location}: BodyFileSHA256 does not match the body file")
        body_plan = {
            "source_path": body_path_text,
            "sha256": actual_digest,
            "content_type": mime_source,
            "delivery": "PUT template enclosure Body after template identity is resolved",
        }

    parent = parent_operations[client_key]
    if parent["operation"] == "CREATE":
        parent_key = str(parent["payload"]["AlertCode"])
    else:
        parent_key = str(parent["matched_identity"]["canonical_key"])
    encoded_parent_key = quote(parent_key, safe="")
    if operation == "CREATE":
        if row.get("TemplateId") not in {None, ""}:
            raise IntakeError(f"{location}: TemplateId must be blank for CREATE")
        method = "POST"
        relative_uri = f"{RESOURCE}/{encoded_parent_key}/child/templates"
    else:
        template_key = str(row.get("TemplateCode") or "").strip()
        method = "PATCH"
        relative_uri = (
            f"{RESOURCE}/{encoded_parent_key}/child/templates/"
            f"{quote(template_key, safe='')}"
        )
        if operation == "DISABLE":
            payload = {"EnabledFlag": False}
    template_operation = {
        "entity": "template",
        "client_alert_key": client_key,
        "template_key": str(row["TemplateKey"]),
        "operation": operation,
        "method": method,
        "relative_uri": relative_uri,
        "payload": payload,
        "parent_canonical_key": parent_key,
        "template_canonical_key": str(payload.get("TemplateCode") or ""),
        "body_attachment_plan": body_plan,
        "source": source,
        "preflight": [
            "parent create must return the approved AlertCode as its canonical key",
            "compatible template describe",
            "live MIME and language lookup validation",
            "immutable work-unit preview and aggregate immediate approval",
        ],
    }
    if has_inline_body:
        if operation != "CREATE":
            raise IntakeError(
                f"{location}: inline Body is supported only for Template CREATE"
            )
        template_item_uri = (
            f"{RESOURCE}/{encoded_parent_key}/child/templates/"
            f"{quote(str(payload['TemplateCode']), safe='')}"
        )
        template_operation["body_patch_operation"] = {
            "entity": "template_body",
            "client_alert_key": client_key,
            "template_key": str(row["TemplateKey"]),
            "operation": "CREATE",
            "method": "PATCH",
            "relative_uri": template_item_uri,
            "payload": {"Body": inline_body},
            "parent_canonical_key": parent_key,
            "template_canonical_key": str(payload["TemplateCode"]),
            "source": source,
            "preflight": [
                "parent and Template creates must return their approved canonical keys",
                "plain-text Body PATCH payload must match the immutable preview",
                "PATCH HTTP 200 acceptance is the version 0.8.0 verification contract",
                "independent Body read-back is deferred",
            ],
        }
    else:
        template_operation["body_patch_operation"] = None
    return template_operation


def build_recipient_operation(
    row: dict[str, Any],
    parent_template_operations: dict[tuple[str, str], dict[str, Any]],
) -> dict[str, Any]:
    source = row["_source"]
    location = f"{source['sheet']}!{source['row']}"
    for field in (
        "ClientAlertKey",
        "TemplateKey",
        "RecipientKey",
        "Operation",
        "RecipientCode",
        "ChannelType",
        "ValueExpression",
    ):
        required(row, field, location=location)

    client_key = str(row["ClientAlertKey"])
    template_key = str(row["TemplateKey"])
    parent_key = (client_key, template_key)
    parent = parent_template_operations.get(parent_key)
    if parent is None:
        raise IntakeError(
            f"{location}: no included template for ClientAlertKey {client_key} "
            f"and TemplateKey {template_key}"
        )

    operation = str(row["Operation"]).strip().upper()
    if operation != "CREATE":
        raise IntakeError(
            f"{location}: basic reader supports CREATE recipients only"
        )
    if parent["operation"] != "CREATE":
        raise IntakeError(
            f"{location}: recipient CREATE currently requires an included CREATE template"
        )
    if row.get("RecipientId") not in {None, ""}:
        raise IntakeError(f"{location}: RecipientId must be blank for CREATE")

    payload = {field: row.get(field) for field in RECIPIENT_FIELDS}
    if len(str(payload["RecipientCode"])) > 32:
        raise IntakeError(f"{location}: RecipientCode exceeds 32 characters")
    if len(str(payload["ChannelType"])) > 30:
        raise IntakeError(f"{location}: ChannelType exceeds 30 characters")
    if len(str(payload["ValueExpression"])) > 2000:
        raise IntakeError(f"{location}: ValueExpression exceeds 2000 characters")

    parent_canonical_key = str(parent["parent_canonical_key"])
    template_canonical_key = str(parent["template_canonical_key"])
    return {
        "entity": "template_recipient",
        "client_alert_key": client_key,
        "template_key": template_key,
        "recipient_key": str(row["RecipientKey"]),
        "operation": operation,
        "method": "POST",
        "relative_uri": (
            f"{RESOURCE}/{quote(parent_canonical_key, safe='')}/child/templates/"
            f"{quote(template_canonical_key, safe='')}/child/recipients"
        ),
        "payload": payload,
        "source": source,
        "preflight": [
            "parent and template creates must return their approved canonical keys",
            "compatible recipient describe",
            "live ORA_HRC_ALERT_CHANNEL_TYPE validation",
            "ValueExpression required and maximum 2,000 characters; live validation deferred",
            "immutable work-unit preview and aggregate immediate approval",
        ],
    }


def operation_verification_uri(operation: dict[str, Any]) -> str:
    if operation["entity"] == "alert":
        if operation["method"] == "POST":
            key = quote(str(operation["payload"]["AlertCode"]), safe="")
            return f"{RESOURCE}/{key}"
        return str(operation["relative_uri"])
    if operation["entity"] == "template":
        if operation["method"] == "POST":
            key = quote(str(operation["payload"]["TemplateCode"]), safe="")
            return f"{operation['relative_uri']}/{key}"
        return str(operation["relative_uri"])
    if operation["entity"] == "template_recipient":
        key = quote(str(operation["payload"]["RecipientCode"]), safe="")
        return f"{operation['relative_uri']}/{key}"
    raise IntakeError(f"Unsupported work-unit entity {operation['entity']!r}")


def operation_verification(operation: dict[str, Any], target_environment: str | None) -> dict[str, Any]:
    if operation["entity"] == "template_body":
        return {
            "mode": "mutation_response",
            "accepted_http_statuses": [200],
            "independent_read_back": False,
            "deferred_enhancement": (
                "Verify stored Body content through a supported Body/enclosure read route"
            ),
        }
    verification_relative_uri = operation_verification_uri(operation)
    return {
        "method": "GET",
        "uri": (
            f"{target_environment}{verification_relative_uri}"
            if target_environment
            else verification_relative_uri
        ),
        "compare_fields": sorted(operation["payload"]),
    }


def order_operations(operations: list[dict[str, Any]]) -> list[dict[str, Any]]:
    ordered: list[dict[str, Any]] = []
    for alert in (item for item in operations if item["entity"] == "alert"):
        ordered.append(alert)
        templates = [
            item
            for item in operations
            if item["entity"] == "template"
            and item["client_alert_key"] == alert["client_alert_key"]
        ]
        for template in templates:
            ordered.append(template)
            ordered.extend(
                item
                for item in operations
                if item["entity"] == "template_body"
                and item["client_alert_key"] == alert["client_alert_key"]
                and item["template_key"] == template["template_key"]
            )
            ordered.extend(
                item
                for item in operations
                if item["entity"] == "template_recipient"
                and item["client_alert_key"] == alert["client_alert_key"]
                and item["template_key"] == template["template_key"]
            )
    return ordered


def operation_effect(operation: dict[str, Any]) -> str:
    labels = {
        "alert": "Create the parent Alert",
        "template": "Create a no-Body Alert Template",
        "template_body": "Set the Template Body by PATCH",
        "template_recipient": "Create a Template Recipient",
    }
    return labels[operation["entity"]]


def build_create_work_units(
    operations: list[dict[str, Any]],
    *,
    source_sha256: str,
    target_environment: str | None,
    validation_mode: str = "simple",
) -> list[dict[str, Any]]:
    work_units: list[dict[str, Any]] = []
    parents = [item for item in operations if item["entity"] == "alert"]
    for parent in parents:
        client_key = parent["client_alert_key"]
        unit_operations = [
            item for item in operations if item["client_alert_key"] == client_key
        ]
        eligible = all(item["operation"] == "CREATE" for item in unit_operations)
        eligible = eligible and all(
            item.get("body_attachment_plan") is None
            for item in unit_operations
            if item["entity"] == "template"
        )
        if not eligible:
            continue

        approved_operations: list[dict[str, Any]] = []
        for sequence, operation in enumerate(unit_operations, start=1):
            relative_uri = str(operation["relative_uri"])
            approved_operations.append(
                {
                    "sequence": sequence,
                    "entity": operation["entity"],
                    "method": operation["method"],
                    "uri": (
                        f"{target_environment}{relative_uri}"
                        if target_environment
                        else relative_uri
                    ),
                    "payload": operation["payload"],
                    "payload_sha256": canonical_sha256(operation["payload"]),
                    "effect": operation_effect(operation),
                    "verification": operation_verification(
                        operation, target_environment
                    ),
                }
            )

        approval_document = {
            "contract_version": CONTRACT_VERSION,
            "execution_mode": "sequential",
            "validation_mode": validation_mode,
            "validation_policy": alert_validation.policy(validation_mode),
            "approval_scope": "one_alert_hierarchy_create",
            "target_environment": target_environment,
            "source_artifact_sha256": source_sha256,
            "execution_policy": {
                "ordered": True,
                "verify_after_each_mutation": True,
                "stop_on_first_failure": True,
                "automatic_write_retry": False,
                "automatic_rollback": False,
                "canonical_key_mismatch_stops_unit": True,
            },
            "operations": approved_operations,
        }
        if any(item["entity"] == "template_body" for item in unit_operations):
            approval_document["execution_policy"]["body_patch_verification"] = (
                "accepted_http_200_response"
            )
        fingerprint = (
            canonical_sha256(approval_document) if target_environment else None
        )
        work_units.append(
            {
                "client_alert_key": client_key,
                "approval_binding_complete": target_environment is not None,
                "approval_blocker": (
                    ("Deep provider preflight gates must pass before approval" if validation_mode == "deep" else None)
                    if target_environment
                    else "target_environment is required before approval"
                ),
                "approval_document": approval_document,
                "work_unit_sha256": fingerprint,
                "approval_phrase": (
                    f"Confirm Alert Work Unit {fingerprint}"
                    if fingerprint
                    else None
                ),
            }
        )
    return work_units


def build_plan(args: argparse.Namespace) -> dict[str, Any]:
    validation_mode = getattr(args, "validation_mode", "simple")
    try:
        validation_policy = alert_validation.policy(validation_mode)
    except ValueError as exc:
        raise IntakeError(str(exc)) from exc
    if validation_mode == "simple" and args.mime_type_map:
        raise IntakeError("Simple mode uses literal MimeType codes; mappings require deep mode")
    workbook = Path(args.workbook).expanduser().resolve()
    values_wb, formulas_wb = open_workbook(workbook)
    try:
        alerts = read_rows(values_wb, formulas_wb, "Alerts")
        templates = read_rows(values_wb, formulas_wb, "Alert Templates")
        recipients = (
            read_rows(values_wb, formulas_wb, "Template Recipients")
            if "Template Recipients" in values_wb.sheetnames
            else []
        )
        for optional_sheet in (
            "Template Tokens",
            "Template Translations",
            "Alert Tokens",
            "Alert Translations",
        ):
            if optional_sheet in values_wb.sheetnames and read_rows(
                values_wb, formulas_wb, optional_sheet
            ):
                raise IntakeError(
                    f"{optional_sheet}: included rows are not supported by the basic reader"
                )
    finally:
        values_wb.close()
        formulas_wb.close()

    if not alerts:
        raise IntakeError("No included Alerts rows were found")
    client_keys = [str(row.get("ClientAlertKey") or "") for row in alerts]
    if len(client_keys) != len(set(client_keys)):
        raise IntakeError("Included Alerts rows contain duplicate ClientAlertKey values")

    current_alerts = load_current_alerts(
        Path(args.current_state).resolve() if args.current_state else None
    )
    operations: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    parent_operations: dict[str, dict[str, Any]] = {}
    for row in alerts:
        try:
            operation = build_alert_operation(row, current_alerts)
        except IntakeError as exc:
            errors.append(
                {
                    "stage": "alert_intake",
                    "source": row.get("_source"),
                    "failure_class": "input_validation",
                    "message": str(exc),
                }
            )
        else:
            operations.append(operation)
            parent_operations[operation["client_alert_key"]] = operation

    mime_map = parse_mapping(args.mime_type_map)
    seen_templates: set[tuple[str, str]] = set()
    parent_template_operations: dict[tuple[str, str], dict[str, Any]] = {}
    for row in templates:
        key = (str(row.get("ClientAlertKey") or ""), str(row.get("TemplateKey") or ""))
        if key in seen_templates:
            raise IntakeError(f"Alert Templates contains duplicate key {key}")
        seen_templates.add(key)
        row["_workbook_dir"] = str(workbook.parent)
        try:
            template_operation = build_template_operation(
                row, parent_operations, mime_map, validation_mode
            )
        except IntakeError as exc:
            errors.append(
                {
                    "stage": "template_intake",
                    "source": row.get("_source"),
                    "failure_class": "input_validation",
                    "message": str(exc),
                }
            )
        else:
            body_patch_operation = template_operation.pop(
                "body_patch_operation", None
            )
            operations.append(template_operation)
            parent_template_operations[key] = template_operation
            if body_patch_operation is not None:
                operations.append(body_patch_operation)

    seen_recipients: set[tuple[str, str, str]] = set()
    for row in recipients:
        key = (
            str(row.get("ClientAlertKey") or ""),
            str(row.get("TemplateKey") or ""),
            str(row.get("RecipientKey") or ""),
        )
        if key in seen_recipients:
            raise IntakeError(f"Template Recipients contains duplicate key {key}")
        seen_recipients.add(key)
        try:
            operations.append(
                build_recipient_operation(row, parent_template_operations)
            )
        except IntakeError as exc:
            errors.append(
                {
                    "stage": "template_recipient_intake",
                    "source": row.get("_source"),
                    "failure_class": "input_validation",
                    "message": str(exc),
                }
            )

    operations = order_operations(operations)
    source_sha256 = sha256_file(workbook)
    target_environment = normalize_target_environment(
        getattr(args, "target_environment", None)
    )
    execution_mode = getattr(args, "execution_mode", "sequential")
    if execution_mode not in {"sequential", "batch-create"}:
        raise IntakeError("Unsupported execution mode")
    for operation in operations:
        if validation_mode == "simple" and operation["operation"] == "CREATE":
            operation["preflight"] = [
                "local mandatory fields, hierarchy, types and payload checks",
                "exact environment and immutable approval binding",
                "Fusion validates business values; no live preflight lookups",
            ]
    if execution_mode == "batch-create":
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        import alert_batch

        units = []
        if not errors:
            try:
                units = [alert_batch.build_work_unit(
                    operations, source_sha256=source_sha256,
                    target_environment=target_environment,
                    validation_mode=validation_mode,
                    attempt_id=getattr(args, "attempt_id", None),
                    recovery_context=getattr(args, "recovery_context", "initial"),
                    project_id_value=getattr(args, "project_id", None),
                )]
            except ValueError as exc:
                errors.append({"stage": "batch_intake", "failure_class": "input_validation", "message": str(exc)})
        return {
            "status": "validation_failed" if errors else "preview_candidate",
            "contract_version": alert_batch.VERSION,
            "execution_mode": execution_mode,
            "validation_mode": validation_mode,
            "validation_policy": validation_policy,
            "preflight_requests": alert_validation.preflight_requests(validation_mode, operations, execution_mode),
            "reader": f"alert-foundation-batch-xlsx/openpyxl/{alert_batch.VERSION}",
            "source": {"workbook": str(workbook), "sha256": source_sha256},
            "approval_mode": "one_approval_for_entire_batch_create",
            "work_units": units,
            "errors": errors,
            "warnings": [
                "Simple mode needs only local gates before approval; deep mode requires its declared live preflight checks.",
                "One POST contains every part; server atomicity and stop-on-first-part-error are unverified.",
                "Inline Body is encoded once as UTF-8 Base64; stored Body verification remains deferred.",
                "HTTP 200 alone is not success; all parts and independent read-backs must be checked.",
                "No automatic write retry or rollback; no Fusion request was performed.",
            ],
        }
    work_units = (
        build_create_work_units(
            operations,
            source_sha256=source_sha256,
            target_environment=target_environment,
            validation_mode=validation_mode,
        )
        if not errors
        else []
    )
    return {
        "status": "validation_failed" if errors else "preview_candidate",
        "contract_version": CONTRACT_VERSION,
        "execution_mode": execution_mode,
        "validation_mode": validation_mode,
        "validation_policy": validation_policy,
        "preflight_requests": alert_validation.preflight_requests(validation_mode, operations, execution_mode),
        "reader": f"alert-foundation-basic-xlsx/openpyxl/{CONTRACT_VERSION}",
        "source": {
            "workbook": str(workbook),
            "sha256": source_sha256,
        },
        "operations": operations,
        "approval_mode": "one_approval_per_alert_hierarchy_create_work_unit",
        "work_units": work_units,
        "errors": errors,
        "warnings": [
            "Workbook validation status and approval cells were ignored.",
            "Simple CREATE uses local gates only; deep mode requires live preflight. Existing-item operations still require fresh exact identity in either mode.",
            "Live Template Recipient expression validation is deferred in provider version 0.8.0.",
            "Inline plain-text Template Body uses a separate PATCH; HTTP 200 acceptance is treated as success and independent Body read-back is deferred.",
            "One aggregate approval covers only the exact ordered operations and fingerprint in one eligible Alert hierarchy work unit.",
            "Any changed environment, URI, method, payload, order, or verification contract invalidates that approval.",
            "This script performs no Fusion request.",
        ],
    }


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(
        description="Read an Alert Foundation XLSX intake and emit a REST payload plan."
    )
    result.add_argument("workbook", help="Path to the Alert Foundation .xlsx workbook")
    result.add_argument("--validation-mode", choices=alert_validation.MODES, default="simple",
                        help="simple (default): local checks; deep: additional declared live preflight")
    result.add_argument("--execution-mode", choices=("sequential", "batch-create"), default="sequential",
                        help="Opt-in batch hierarchy creation; sequential remains the default")
    result.add_argument(
        "--current-state",
        help="Read-only Fusion alert JSON used to calculate UPDATE/DISABLE PATCH fields",
    )
    result.add_argument(
        "--mime-type-map",
        action="append",
        default=[],
        metavar="SOURCE=REST_CODE",
        help="Live-validated MIME mapping, for example text/html=ORA_H",
    )
    result.add_argument(
        "--target-environment",
        help=(
            "HTTPS Fusion origin used to bind exact work-unit URIs and the "
            "aggregate approval fingerprint"
        ),
    )
    result.add_argument("--attempt-id", help="UUID bound into this one-time batch approval (generated when omitted)")
    result.add_argument("--recovery-context", default="initial", help="Non-secret context identifying this initial or recovery attempt")
    result.add_argument("--project-id", help="Active FSCR project ID required for batch-create receipt ownership")
    result.add_argument("--output", help="Optional JSON output path")
    return result


def main() -> int:
    args = parser().parse_args()
    try:
        result = build_plan(args)
        exit_code = 0 if result["status"] == "preview_candidate" else 2
    except (IntakeError, RuntimeError) as exc:
        result = {
            "status": "validation_failed",
            "failure_class": "input_validation"
            if isinstance(exc, IntakeError)
            else "runtime_capability_unavailable",
            "message": str(exc),
            "operations": [],
        }
        exit_code = 2
    rendered = json.dumps(result, indent=2, ensure_ascii=False) + "\n"
    if args.output:
        Path(args.output).write_text(rendered, encoding="utf-8")
    else:
        sys.stdout.write(rendered)
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
