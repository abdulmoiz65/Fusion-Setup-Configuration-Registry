#!/usr/bin/env python3
"""Build a review-only task FSM ZIP from the conversion-policy intake workbook.

This is a local candidate builder. It never contacts Fusion or approves an import.
Run with fresh, same-environment task template and populated export ZIPs.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
import re
import sys
from xml.etree import ElementTree as ET
from zipfile import BadZipFile, ZIP_DEFLATED, ZipFile


TASK = "EXM_MANAGE_CONVERSION_RATES_AND_POLICIES"
PARENT = "EXM_EXCHANGE_RATE_OPTION.csv"
CHILD = "EXM_EXCH_RATE_TOLERANCE.csv"
METADATA = "ASM_SETUP_CSV_METADATA.xml"
AUDIT = {"CreationDate", "CreatedBy", "LastUpdateLogin", "LastUpdateDate", "LastUpdatedBy"}
PARENT_FIELDS = {
    "Conversion Rate Type": "UserConversionType",
    "Default Conversion Rate?": "DfltExchangeRateFlag",
    "Warning Tolerance": "WarningTolerance",
    "Error Tolerance": "ErrorTolerance",
    "Display Warning to User?": "DispWarningToUser",
}
CHILD_FIELDS = {key: value for key, value in PARENT_FIELDS.items() if key != "Conversion Rate Type"}
PARENT_HEAD = ("Business Unit*", *PARENT_FIELDS)
CHILD_HEAD = ("Business Unit*", "Currency Code*", *CHILD_FIELDS)
CHILD_BU = "EXM_EXCHANGE_RATE_OPTION.BuName"


class IntakeError(ValueError):
    pass


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        raise IntakeError("Boolean cells must use the workbook's Yes or No text")
    return str(value).strip()


def _rows(path: Path, sheet_name: str, header: tuple[str, ...]):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise IntakeError("openpyxl is required to read the intake workbook") from exc
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if set(book.sheetnames) != {"Conversion Rate Policies", "Currency Tolerances"}:
            raise IntakeError("Workbook must contain exactly the two documented intake sheets")
        sheet = book[sheet_name]
        if sheet.sheet_state != "visible":
            raise IntakeError(f"{sheet_name}: sheet must be visible")
        actual = tuple(_cell(sheet.cell(4, index).value) for index in range(1, len(header) + 1))
        if actual != header:
            raise IntakeError(f"{sheet_name}: row 4 headers do not match the intake contract")
        result = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=5), start=5):
            values = []
            for cell in row:
                if cell.data_type == "f":
                    raise IntakeError(f"{sheet_name}!{cell.coordinate}: formulas are not accepted")
                values.append(_cell(cell.value))
            if any(values[len(header):]):
                raise IntakeError(f"{sheet_name}!{row_number}: unexpected value after the last input column")
            values = (values + [""] * len(header))[:len(header)]
            if any(values):
                if not values[0] or (sheet_name == "Currency Tolerances" and not values[1]):
                    raise IntakeError(f"{sheet_name}!{row_number}: required identity is blank")
                result.append((row_number, dict(zip(header, values))))
        return result
    finally:
        book.close()


def _package(path: Path, environment: str, max_age_hours: float):
    with ZipFile(path) as archive:
        names = archive.namelist()
        if set(names) != {PARENT, CHILD, METADATA} or len(names) != 3:
            raise IntakeError(f"{path.name}: expected exact task metadata and two CSV members")
        xml = archive.read(METADATA)
        root = ET.fromstring(xml)
        values = {element.tag.split("}")[-1]: (element.text or "").strip() for element in root.iter()}
        if values.get("ObjectShortName") != TASK or values.get("ProcessEnv") != environment:
            raise IntakeError(f"{path.name}: task or environment does not match")
        date_text = values.get("ProcessDate", "")
        try:
            issued = datetime.fromisoformat(date_text.replace("Z", "+00:00"))
        except ValueError as exc:
            raise IntakeError(f"{path.name}: invalid ProcessDate") from exc
        # A copied old export must not become fresh merely through a new mtime.
        # The task export timestamps used by this workflow are interpreted as UTC.
        if issued.tzinfo is None:
            issued = issued.replace(tzinfo=timezone.utc)
        age = (datetime.now(timezone.utc) - issued).total_seconds() / 3600
        if age < -1 or age > max_age_hours:
            raise IntakeError(f"{path.name}: package is not fresh (age {age:.1f} hours)")
        tables = {}
        for member in (PARENT, CHILD):
            content = archive.read(member).decode("utf-8-sig")
            parsed = list(csv.reader(StringIO(content, newline="")))
            if not parsed:
                raise IntakeError(f"{path.name}: {member} has no header")
            head = parsed[0]
            if len(head) != 11 or len(set(head)) != 11:
                raise IntakeError(f"{path.name}: {member} has an invalid header")
            if member == PARENT and set(head) != {"BuName", *PARENT_FIELDS.values(), *AUDIT}:
                raise IntakeError(f"{path.name}: parent fields differ from the reviewed task shape")
            if member == CHILD and set(head) != {CHILD_BU, "CurrencyCode", *CHILD_FIELDS.values(), *AUDIT}:
                raise IntakeError(f"{path.name}: child fields differ from the reviewed task shape")
            if any(len(row) != len(head) for row in parsed[1:]):
                raise IntakeError(f"{path.name}: {member} contains a short or long record")
            tables[member] = (head, [dict(zip(head, row)) for row in parsed[1:]])
        return xml, tables


def _normalize(label: str, raw: str, location: str) -> str:
    if not raw:
        return ""
    if label in {"Default Conversion Rate?", "Display Warning to User?"}:
        if raw.lower() not in {"yes", "no"}:
            raise IntakeError(f"{location}: {label} must be Yes or No")
        return "Y" if raw.lower() == "yes" else "N"
    if label in {"Error Tolerance", "Warning Tolerance"}:
        try:
            number = Decimal(raw)
        except InvalidOperation as exc:
            raise IntakeError(f"{location}: {label} must be numeric") from exc
        if not number.is_finite() or number < 0 or "%" in raw:
            raise IntakeError(f"{location}: {label} must be a nonnegative percentage-point number")
        return format(number, "f")
    if raw[0] in "=+-@" or "\n" in raw or "\r" in raw:
        raise IntakeError(f"{location}: unsafe text value")
    return raw


def _merge(input_row, mapping, existing, location):
    result = dict(existing or {})
    changed = False
    for label, field in mapping.items():
        supplied = _normalize(label, input_row[label], location)
        if supplied:
            changed |= result.get(field) != supplied
            result[field] = supplied
        elif existing is None:
            # No source-backed create defaults are established for this task.
            raise IntakeError(f"{location}: {label} is required for a new row")
    # The live route accepts complete merged business rows only. Reject an
    # unsupported baseline rather than serializing an accidental null request.
    for label, field in mapping.items():
        if not result.get(field):
            raise IntakeError(f"{location}: {label} needs an explicit value for the complete-row contract")
        raw = result[field]
        if label in {"Default Conversion Rate?", "Display Warning to User?"}:
            if raw not in {"Y", "N"}:
                raise IntakeError(f"{location}: {label} has an unsupported baseline flag")
        else:
            _normalize(label, raw, location)
    if Decimal(result["ErrorTolerance"]) <= Decimal(result["WarningTolerance"]):
        raise IntakeError(f"{location}: Error Tolerance must exceed Warning Tolerance")
    for field in AUDIT:
        result[field] = ""
    return result, changed


def build(workbook: Path, template: Path, current_export: Path, environment: str,
          output: Path, max_age_hours: float = 24, dry_run: bool = False,
          *, _memory_archive: BytesIO | None = None) -> dict:
    if not Decimal(str(max_age_hours)).is_finite() or max_age_hours <= 0:
        raise IntakeError("max-age-hours must be positive and finite")
    if _memory_archive is None and output.exists():
        raise IntakeError(f"Output already exists: {output}")
    if not environment or not re.fullmatch(r"[A-Za-z0-9_-]+", environment):
        raise IntakeError("An exact environment identifier is required")
    parent_input = _rows(workbook, "Conversion Rate Policies", PARENT_HEAD)
    child_input = _rows(workbook, "Currency Tolerances", CHILD_HEAD)
    if not parent_input and not child_input:
        raise IntakeError("Workbook has no policy or currency-tolerance rows")
    template_xml, template_tables = _package(template, environment, max_age_hours)
    _, export_tables = _package(current_export, environment, max_age_hours)
    if any(rows for _, rows in template_tables.values()):
        raise IntakeError("Template CSV members must contain headers only")
    if any(template_tables[name][0] != export_tables[name][0] for name in (PARENT, CHILD)):
        raise IntakeError("Template and export CSV headers differ")
    parents = {}
    for row in export_tables[PARENT][1]:
        key = row["BuName"]
        if not key or key in parents:
            raise IntakeError(f"Current export has a missing or duplicate business unit: {key!r}")
        parents[key] = row
    children = {}
    for row in export_tables[CHILD][1]:
        key = (row[CHILD_BU], row["CurrencyCode"])
        if not all(key) or key in children or key[0] not in parents:
            raise IntakeError(f"Current export has an invalid or duplicate child: {key!r}")
        children[key] = row
    planned_parents, planned_children = {}, {}
    changes = []
    for number, data in parent_input:
        key = _normalize("Business Unit*", data["Business Unit*"], f"Conversion Rate Policies!{number}")
        if key in planned_parents:
            raise IntakeError(f"Conversion Rate Policies!{number}: duplicate business unit")
        location = f"Conversion Rate Policies!{number}"
        merged, changed = _merge(data, PARENT_FIELDS, parents.get(key), location)
        merged["BuName"] = key
        planned_parents[key] = merged
        if changed or key not in parents:
            changes.append(location)
    for number, data in child_input:
        bu = _normalize("Business Unit*", data["Business Unit*"], f"Currency Tolerances!{number}")
        currency = data["Currency Code*"].upper()
        location = f"Currency Tolerances!{number}"
        if not re.fullmatch(r"[A-Z]{3}", currency):
            raise IntakeError(f"{location}: Currency Code must be three letters")
        key = bu, currency
        if key in planned_children:
            raise IntakeError(f"{location}: duplicate business unit and currency")
        if bu not in planned_parents:
            if bu not in parents:
                raise IntakeError(f"{location}: business unit has no parent policy")
            planned_parents[bu], _ = _merge(
                {field: "" for field in PARENT_FIELDS}, PARENT_FIELDS, parents[bu], location
            )
        merged, changed = _merge(data, CHILD_FIELDS, children.get(key), location)
        merged[CHILD_BU], merged["CurrencyCode"] = key
        planned_children[key] = merged
        if changed or key not in children:
            changes.append(location)
    if not changes:
        raise IntakeError("Workbook produces no changes against the current export")
    # Preserve every observed child of a touched parent in the review candidate.
    # Header-only child CSV preservation was observed; absent members/nodes and
    # partially omitted siblings remain unproved. Retain conservative carry-forward.
    for key, existing in children.items():
        if key[0] in planned_parents and key not in planned_children:
            planned_children[key], _ = _merge(
                {field: "" for field in CHILD_FIELDS}, CHILD_FIELDS, existing, f"export child {key}"
            )
    field_changes = []
    for member, planned, existing, mapping in (
        (PARENT, planned_parents, parents, PARENT_FIELDS),
        (CHILD, planned_children, children, CHILD_FIELDS),
    ):
        for key, row in planned.items():
            previous = existing.get(key)
            fields = [field for field in mapping.values()
                      if previous is None or row.get(field, "") != previous.get(field, "")]
            if fields:
                field_changes.append({"member": member, "identity": key,
                                      "operation": "create" if previous is None else "update",
                                      "changed_fields": fields})
    result = {"status": "REVIEW_ONLY_NOT_APPROVED", "task_code": TASK,
              "environment": environment, "dry_run": dry_run,
              "parent_rows": len(planned_parents), "child_rows": len(planned_children),
              "changed_input_rows": changes, "field_changes": field_changes,
              "blockers": ["Runtime destination reference checks and exact-candidate Task Hook validation remain required",
                           "Clearing, default-dependent creates, rekeying and deletion are unsupported",
                           "Task readiness, reference checks and Data Loader approval are required"]}
    if dry_run:
        return result
    memory = _memory_archive if _memory_archive is not None else BytesIO()
    with ZipFile(memory, "w", compression=ZIP_DEFLATED) as archive:
        for name, planned in ((PARENT, planned_parents.values()), (CHILD, planned_children.values())):
            head = template_tables[name][0]
            content = StringIO(newline="")
            writer = csv.writer(content, lineterminator="\r\n")
            writer.writerow(head)
            writer.writerows([row.get(field, "") for field in head] for row in planned)
            archive.writestr(name, content.getvalue().encode("utf-8"))
        archive.writestr(METADATA, template_xml)
    if _memory_archive is None:
        output.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
        # This candidate may later be bound to core approval. Create exclusively
        # with private permissions, including when a dangling symlink exists.
        descriptor = os.open(output, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        try:
            with os.fdopen(descriptor, "wb") as target:
                target.write(memory.getvalue())
        except BaseException:
            output.unlink()
            raise
    return result


def build_in_memory(workbook: Path, template: Path, current_export: Path,
                    environment: str, max_age_hours: float = 24) -> tuple[dict, bytes]:
    """Use the same intake merge for local checks without persisting a ZIP.

    This produces review-only bytes; the caller still owns Task Hook validation,
    reference preflight and a later candidate-bound Data Loader approval.
    """
    memory = BytesIO()
    # The public disk builder's overwrite gate is irrelevant for this sink.
    result = build(workbook, template, current_export, environment,
                   Path('__memory_only_fsm_candidate__.zip'), max_age_hours,
                   _memory_archive=memory)
    return result, memory.getvalue()


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--current-export", required=True, type=Path)
    parser.add_argument("--environment", required=True)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--max-age-hours", type=float, default=24)
    parser.add_argument("--dry-run", action="store_true", help="Validate and preview field changes without writing a ZIP")
    args = parser.parse_args(argv)
    try:
        if args.max_age_hours <= 0:
            raise IntakeError("max-age-hours must be positive")
        result = build(args.workbook, args.template, args.current_export,
                       args.environment, args.output, args.max_age_hours, args.dry_run)
    except (ValueError, OSError, BadZipFile, ET.ParseError, csv.Error) as exc:
        print(json.dumps({"status": "FAILED", "stage": "candidate_preparation",
                          "failure_class": "input_validation", "message": str(exc)}), file=sys.stderr)
        return 2
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
